#!/usr/bin/env python3
"""Validate structural invariants of a generated DOCX.

This complements visual inspection. It does not prove that a page looks good;
it catches broken packages, missing media, missing field targets, and table
border/alignment invariants that are easy to miss in a screenshot.
"""

from __future__ import annotations

import argparse
import json
import posixpath
import re
import sys
import xml.etree.ElementTree as ET
from pathlib import Path
from zipfile import BadZipFile, ZipFile

try:
    from docx import Document
    from docx.oxml.ns import qn
except ImportError as exc:  # pragma: no cover - environment-dependent
    print(f"ERROR: python-docx is required: {exc}", file=sys.stderr)
    raise SystemExit(2)


NS = {
    "w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main",
    "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
    "pr": "http://schemas.openxmlformats.org/package/2006/relationships",
}


def local_text(element: ET.Element) -> str:
    return "".join(element.itertext())


def field_instructions(root: ET.Element) -> list[str]:
    return [local_text(node).strip() for node in root.findall(".//w:instrText", NS) if local_text(node).strip()]


def bookmark_sets(root: ET.Element) -> tuple[set[str], set[str]]:
    starts = {
        node.get(f"{{{NS['w']}}}name", "")
        for node in root.findall(".//w:bookmarkStart", NS)
        if node.get(f"{{{NS['w']}}}name")
    }
    ends = {
        node.get(f"{{{NS['w']}}}id", "")
        for node in root.findall(".//w:bookmarkEnd", NS)
        if node.get(f"{{{NS['w']}}}id")
    }
    # The names/IDs are checked separately because Word stores the name only
    # on bookmarkStart and the numeric ID on both start/end nodes.
    start_ids = {
        node.get(f"{{{NS['w']}}}id", "")
        for node in root.findall(".//w:bookmarkStart", NS)
        if node.get(f"{{{NS['w']}}}id")
    }
    return starts, start_ids - ends


def relationship_issues(zf: ZipFile, document_root: ET.Element) -> list[str]:
    issues: list[str] = []
    rels_name = "word/_rels/document.xml.rels"
    if rels_name not in zf.namelist():
        return [f"missing {rels_name}"]
    rels_root = ET.fromstring(zf.read(rels_name))
    relations = {
        rel.get("Id"): rel
        for rel in rels_root.findall("pr:Relationship", NS)
        if rel.get("Id")
    }
    for node in document_root.findall(".//*[@r:embed]", NS):
        rid = node.get(f"{{{NS['r']}}}embed")
        rel = relations.get(rid)
        if rel is None:
            issues.append(f"missing relationship target for {rid}")
            continue
        if rel.get("TargetMode") == "External":
            continue
        target = posixpath.normpath(posixpath.join("word", rel.get("Target", "")))
        if target not in zf.namelist():
            issues.append(f"missing package part {target} for {rid}")
    return issues


def edge_value(parent: ET.Element | None, edge: str) -> tuple[str | None, str | None]:
    if parent is None:
        return None, None
    node = parent.find(f"w:{edge}", NS)
    if node is None:
        return None, None
    return node.get(f"{{{NS['w']}}}val"), node.get(f"{{{NS['w']}}}sz")


def table_report(doc, require_three_line: bool) -> tuple[list[dict], list[str]]:
    reports: list[dict] = []
    issues: list[str] = []
    expected = {
        "top": ("single", "12"),
        "bottom": ("single", "12"),
        "left": ("nil", "0"),
        "right": ("nil", "0"),
        "insideH": ("nil", "0"),
        "insideV": ("nil", "0"),
    }
    for index, table in enumerate(doc.tables, 1):
        tbl_borders = table._tbl.tblPr.first_child_found_in("w:tblBorders")
        edges = {edge: edge_value(tbl_borders, edge) for edge in expected}
        header_repeat = table.rows[0]._tr.get_or_add_trPr().find(qn("w:tblHeader")) is not None
        row_split_flags = []
        for row in table.rows:
            trpr = row._tr.get_or_add_trPr()
            row_split_flags.append(trpr.find(qn("w:cantSplit")) is not None)
        reports.append({
            "index": index,
            "rows": len(table.rows),
            "columns": len(table.columns),
            "borders": edges,
            "repeated_header": header_repeat,
            "rows_cant_split": sum(row_split_flags),
        })
        if require_three_line:
            for edge, wanted in expected.items():
                if edges[edge] != wanted:
                    issues.append(f"table {index}: {edge} border is {edges[edge]}, expected {wanted}")
            if not header_repeat:
                issues.append(f"table {index}: first row is not marked as a repeating header")
            for cell_index, cell in enumerate(table.rows[0].cells, 1):
                tc_borders = cell._tc.get_or_add_tcPr().first_child_found_in("w:tcBorders")
                bottom = edge_value(tc_borders, "bottom")
                if bottom != ("single", "8"):
                    issues.append(f"table {index} header cell {cell_index}: bottom border is {bottom}, expected ('single', '8')")
    return reports, issues


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("docx", type=Path)
    parser.add_argument("--three-line", action="store_true", help="require academic three-line table borders")
    parser.add_argument("--strict", action="store_true", help="exit non-zero for warnings as well as errors")
    parser.add_argument("--json-out", type=Path, help="write the report as JSON")
    args = parser.parse_args()

    path = args.docx.resolve()
    errors: list[str] = []
    warnings: list[str] = []
    report: dict[str, object] = {"path": str(path), "errors": errors, "warnings": warnings}

    try:
        with ZipFile(path) as zf:
            zip_test = zf.testzip()
            report["zip_test"] = zip_test
            if zip_test is not None:
                errors.append(f"ZIP CRC failure at {zip_test}")
            document_root = ET.fromstring(zf.read("word/document.xml"))
            errors.extend(relationship_issues(zf, document_root))
            instructions = field_instructions(document_root)
            report["fields"] = {
                "REF": sum(1 for item in instructions if re.search(r"\bREF\b", item, re.I)),
                "SEQ": sum(1 for item in instructions if re.search(r"\bSEQ\b", item, re.I)),
                "HYPERLINK": sum(1 for item in instructions if re.search(r"\bHYPERLINK\b", item, re.I)),
                "instructions": instructions,
            }
            bookmark_names, unmatched_ids = bookmark_sets(document_root)
            report["bookmarks"] = {"names": sorted(bookmark_names), "unmatched_start_ids": sorted(unmatched_ids)}
            if unmatched_ids:
                errors.append(f"unmatched bookmark start IDs: {sorted(unmatched_ids)}")

            ref_targets = []
            for instruction in instructions:
                match = re.search(r"\bREF\s+([A-Za-z0-9_]+)", instruction, re.I)
                if match:
                    ref_targets.append(match.group(1))
            missing_ref_targets = sorted(set(ref_targets) - bookmark_names)
            report["missing_ref_targets"] = missing_ref_targets
            if missing_ref_targets:
                errors.append(f"REF fields point to missing bookmarks: {missing_ref_targets}")
    except (BadZipFile, FileNotFoundError, KeyError, ET.ParseError) as exc:
        errors.append(f"cannot inspect DOCX package: {exc}")
        report["exception"] = str(exc)

    try:
        doc = Document(path)
        report["document"] = {
            "paragraphs": len(doc.paragraphs),
            "tables": len(doc.tables),
            "inline_shapes": len(doc.inline_shapes),
            "table_sizes": [f"{len(table.rows)}x{len(table.columns)}" for table in doc.tables],
        }
        table_reports, table_issues = table_report(doc, args.three_line)
        report["tables"] = table_reports
        errors.extend(table_issues)

        duplicate_heading_candidates = []
        for paragraph in doc.paragraphs:
            text = paragraph.text.strip()
            if re.search(r"^(\d+(?:\.\d+){1,3}\.?\s+){2}", text):
                duplicate_heading_candidates.append(text)
        report["duplicate_heading_candidates"] = duplicate_heading_candidates
        if duplicate_heading_candidates:
            warnings.append("possible duplicated manual/list heading prefixes detected")
    except Exception as exc:  # python-docx has several version-specific exceptions
        errors.append(f"python-docx could not open the file: {exc}")

    if args.json_out:
        args.json_out.parent.mkdir(parents=True, exist_ok=True)
        args.json_out.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")

    print(f"docx: {path}")
    print(f"errors: {len(errors)}; warnings: {len(warnings)}")
    if "document" in report:
        print(f"paragraphs/tables/images: {report['document']['paragraphs']} / {report['document']['tables']} / {report['document']['inline_shapes']}")
    fields = report.get("fields", {})
    if fields:
        print(f"fields REF/SEQ/HYPERLINK: {fields['REF']} / {fields['SEQ']} / {fields['HYPERLINK']}")
    if report.get("missing_ref_targets"):
        print(f"missing REF targets: {report['missing_ref_targets']}")
    if errors:
        print("ERRORS:")
        for item in errors:
            print(f"  - {item}")
    if warnings:
        print("WARNINGS:")
        for item in warnings:
            print(f"  - {item}")
    if args.json_out:
        print(f"json: {args.json_out.resolve()}")

    if errors or (args.strict and warnings):
        return 1
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
